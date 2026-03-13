#!/usr/bin/env python3
import argparse
import base64
import csv
import json
import subprocess
import zipfile
from pathlib import Path


MAX_VISUAL_ASSETS = 6
MAX_TEXT_CHARS = 200000
DEFAULT_MAX_PDF_PAGES = 120
DEFAULT_MAX_SHEETS = 10
DEFAULT_MAX_ROWS_PER_SHEET = 5000
DEFAULT_MAX_COLS = 50
DEFAULT_MAX_CELLS = 100000


def clip_text(text: str, limit: int = MAX_TEXT_CHARS) -> str:
    text = (text or "").replace("\r\n", "\n").replace("\x00", "").strip()
    if len(text) <= limit:
        return text
    return text[:limit] + "..."


def read_text_file(file_path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return file_path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return file_path.read_text(encoding="utf-8", errors="ignore")


def parse_text_like(file_path: Path, **_kwargs):
    text = read_text_file(file_path)
    return {
        "text": clip_text(text),
        "structuredText": clip_text(text),
        "formatSummary": "保留了原始文本顺序。",
        "stats": {},
        "visualAssets": [],
    }


def parse_csv_file(
    file_path: Path,
    max_rows_per_sheet: int = DEFAULT_MAX_ROWS_PER_SHEET,
    max_cols: int = DEFAULT_MAX_COLS,
    max_cells: int = DEFAULT_MAX_CELLS,
    **_kwargs,
):
    rows = []
    cell_count = 0
    max_cols_seen = 0
    with file_path.open("r", encoding="utf-8", errors="ignore", newline="") as handle:
        reader = csv.reader(handle)
        for row in reader:
            values = [str(cell).strip() for cell in row if str(cell).strip()]
            if not values:
                continue
            if max_rows_per_sheet and len(rows) >= max_rows_per_sheet:
                raise RuntimeError(f"工作表行数超过限制，最多支持 {max_rows_per_sheet} 行")
            limited_values = values[:max_cols] if max_cols and max_cols > 0 else values
            max_cols_seen = max(max_cols_seen, len(limited_values))
            cell_count += len(limited_values)
            if max_cells and cell_count > max_cells:
                raise RuntimeError(f"表格总单元格数量超过限制，最多支持 {max_cells} 个")
            rows.append(" | ".join(limited_values))
    text = "\n".join(rows)
    return {
        "text": clip_text(text),
        "structuredText": clip_text(text),
        "formatSummary": f"保留了表格结构，约 {len(rows)} 行数据。",
        "stats": {
            "rowCount": len(rows),
            "sheetCount": 1,
            "cellCount": cell_count,
            "maxCols": max_cols_seen,
        },
        "visualAssets": [],
    }


def parse_pdf(file_path: Path, max_pages: int = DEFAULT_MAX_PDF_PAGES, **_kwargs):
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    if max_pages and len(reader.pages) > max_pages:
        raise RuntimeError(f"PDF 页数超过限制，最多支持 {max_pages} 页")
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        pages.append(f"[第 {index} 页]\n{page_text.strip()}")
    text = "\n\n".join([page for page in pages if page.strip()])
    return {
        "text": clip_text(text),
        "structuredText": clip_text(text),
        "formatSummary": f"按页保留正文顺序，共 {len(reader.pages)} 页。",
        "stats": {
            "pageCount": len(reader.pages),
        },
        "visualAssets": [],
    }


def parse_docx(file_path: Path, **_kwargs):
    from docx import Document

    document = Document(str(file_path))
    paragraphs = [para.text.strip() for para in document.paragraphs if para.text and para.text.strip()]
    text = "\n".join(paragraphs)
    visual_assets = []
    with zipfile.ZipFile(file_path, "r") as archive:
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue
            if len(visual_assets) >= MAX_VISUAL_ASSETS:
                break
            data = archive.read(name)
            if not data:
                continue
            extension = Path(name).suffix.lower().lstrip(".")
            mime_type = {
                "jpg": "image/jpeg",
                "jpeg": "image/jpeg",
                "png": "image/png",
                "gif": "image/gif",
                "webp": "image/webp",
            }.get(extension, "image/png")
            visual_assets.append({
                "name": Path(name).name,
                "label": f"文档图片 {len(visual_assets) + 1}",
                "sourceType": "embedded-image",
                "mimeType": mime_type,
                "dataBase64": base64.b64encode(data).decode("utf-8"),
            })
    return {
        "text": clip_text(text),
        "structuredText": clip_text(text),
        "formatSummary": "保留了段落顺序和常见文档结构。",
        "stats": {},
        "visualAssets": visual_assets,
    }


def parse_doc(file_path: Path, **_kwargs):
    try:
        completed = subprocess.run(
            ["antiword", str(file_path)],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
        )
    except FileNotFoundError as exc:
        raise RuntimeError("当前沙箱不支持旧版 DOC 文件，请先转成 DOCX 再上传") from exc
    text = completed.stdout or ""
    return {
        "text": clip_text(text),
        "structuredText": clip_text(text),
        "formatSummary": "旧版 Word 仅提取了正文文本。",
        "stats": {},
        "visualAssets": [],
    }


def parse_xlsx(
    file_path: Path,
    max_sheets: int = DEFAULT_MAX_SHEETS,
    max_rows_per_sheet: int = DEFAULT_MAX_ROWS_PER_SHEET,
    max_cols: int = DEFAULT_MAX_COLS,
    max_cells: int = DEFAULT_MAX_CELLS,
    **_kwargs,
):
    import openpyxl

    workbook = openpyxl.load_workbook(str(file_path), data_only=True)
    if max_sheets and len(workbook.worksheets) > max_sheets:
        raise RuntimeError(f"工作表数量超过限制，最多支持 {max_sheets} 个")
    sections = []
    row_count = 0
    cell_count = 0
    max_cols_seen = 0
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if not values:
                continue
            if max_rows_per_sheet and len(rows) >= max_rows_per_sheet:
                raise RuntimeError(f"工作表 {sheet.title} 行数超过限制，最多支持 {max_rows_per_sheet} 行")
            limited_values = values[:max_cols] if max_cols and max_cols > 0 else values
            max_cols_seen = max(max_cols_seen, len(limited_values))
            cell_count += len(limited_values)
            if max_cells and cell_count > max_cells:
                raise RuntimeError(f"表格总单元格数量超过限制，最多支持 {max_cells} 个")
            rows.append(" | ".join(limited_values))
        row_count += len(rows)
        if rows:
            sections.append(f"工作表：{sheet.title}\n" + "\n".join(rows))

    visual_assets = []
    with zipfile.ZipFile(file_path, "r") as archive:
        for name in archive.namelist():
            if not name.startswith("xl/media/"):
                continue
            if len(visual_assets) >= MAX_VISUAL_ASSETS:
                break
            data = archive.read(name)
            if not data:
                continue
            extension = Path(name).suffix.lower().lstrip(".")
            mime_type = {
                "jpg": "image/jpeg",
                "jpeg": "image/jpeg",
                "png": "image/png",
                "gif": "image/gif",
                "webp": "image/webp",
            }.get(extension, "image/png")
            visual_assets.append({
                "name": Path(name).name,
                "label": f"工作簿图片 {len(visual_assets) + 1}",
                "sourceType": "embedded-image",
                "mimeType": mime_type,
                "dataBase64": base64.b64encode(data).decode("utf-8"),
            })

    text = "\n\n".join(sections)
    return {
        "text": clip_text(text),
        "structuredText": clip_text(text),
        "formatSummary": f"保留了表格结构，包含 {len(workbook.worksheets)} 个工作表，约 {row_count} 行数据。",
        "stats": {
            "sheetCount": len(workbook.worksheets),
            "rowCount": row_count,
            "cellCount": cell_count,
            "maxCols": max_cols_seen,
        },
        "visualAssets": visual_assets,
    }


def parse_xls(
    file_path: Path,
    max_sheets: int = DEFAULT_MAX_SHEETS,
    max_rows_per_sheet: int = DEFAULT_MAX_ROWS_PER_SHEET,
    max_cols: int = DEFAULT_MAX_COLS,
    max_cells: int = DEFAULT_MAX_CELLS,
    **_kwargs,
):
    import xlrd

    book = xlrd.open_workbook(str(file_path))
    if max_sheets and book.nsheets > max_sheets:
        raise RuntimeError(f"工作表数量超过限制，最多支持 {max_sheets} 个")
    sections = []
    row_count = 0
    cell_count = 0
    max_cols_seen = 0
    for index in range(book.nsheets):
        sheet = book.sheet_by_index(index)
        rows = []
        for row_index in range(sheet.nrows):
            values = [str(sheet.cell_value(row_index, col_index)).strip() for col_index in range(sheet.ncols)]
            values = [value for value in values if value]
            if not values:
                continue
            if max_rows_per_sheet and len(rows) >= max_rows_per_sheet:
                raise RuntimeError(f"工作表 {sheet.name} 行数超过限制，最多支持 {max_rows_per_sheet} 行")
            limited_values = values[:max_cols] if max_cols and max_cols > 0 else values
            max_cols_seen = max(max_cols_seen, len(limited_values))
            cell_count += len(limited_values)
            if max_cells and cell_count > max_cells:
                raise RuntimeError(f"表格总单元格数量超过限制，最多支持 {max_cells} 个")
            rows.append(" | ".join(limited_values))
        row_count += len(rows)
        if rows:
            sections.append(f"工作表：{sheet.name}\n" + "\n".join(rows))
    text = "\n\n".join(sections)
    return {
        "text": clip_text(text),
        "structuredText": clip_text(text),
        "formatSummary": f"保留了表格结构，包含 {book.nsheets} 个工作表，约 {row_count} 行数据。",
        "stats": {
            "sheetCount": book.nsheets,
            "rowCount": row_count,
            "cellCount": cell_count,
            "maxCols": max_cols_seen,
        },
        "visualAssets": [],
    }


def detect_parser(extension: str):
    extension = (extension or "").lower()
    if extension in {"txt", "md", "markdown", "json", "js", "jsx", "ts", "tsx", "py", "css", "html", "xml", "yaml", "yml", "sql", "sh", "log", "ini", "conf"}:
        return parse_text_like
    if extension == "csv":
        return parse_csv_file
    if extension == "pdf":
        return parse_pdf
    if extension == "docx":
        return parse_docx
    if extension == "doc":
        return parse_doc
    if extension == "xlsx":
        return parse_xlsx
    if extension == "xls":
        return parse_xls
    raise RuntimeError(f"暂不支持解析该文件类型：{extension}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--original-name", default="file")
    parser.add_argument("--extension", default="")
    parser.add_argument("--mime-type", default="")
    parser.add_argument("--max-pages", type=int, default=DEFAULT_MAX_PDF_PAGES)
    parser.add_argument("--max-sheets", type=int, default=DEFAULT_MAX_SHEETS)
    parser.add_argument("--max-rows-per-sheet", type=int, default=DEFAULT_MAX_ROWS_PER_SHEET)
    parser.add_argument("--max-cols", type=int, default=DEFAULT_MAX_COLS)
    parser.add_argument("--max-cells", type=int, default=DEFAULT_MAX_CELLS)
    args = parser.parse_args()

    file_path = Path(args.input)
    output_path = Path(args.output)
    extension = args.extension or file_path.suffix.lstrip(".")
    parser_fn = detect_parser(extension)
    parsed = parser_fn(
        file_path,
        max_pages=args.max_pages,
        max_sheets=args.max_sheets,
        max_rows_per_sheet=args.max_rows_per_sheet,
        max_cols=args.max_cols,
        max_cells=args.max_cells,
    )

    output = {
        "originalName": args.original_name,
        "extension": extension,
        "mimeType": args.mime_type,
        "text": parsed.get("text") or "",
        "structuredText": parsed.get("structuredText") or parsed.get("text") or "",
        "formatSummary": parsed.get("formatSummary") or "",
        "stats": parsed.get("stats") or {},
        "visualAssets": parsed.get("visualAssets") or [],
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(output, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()
